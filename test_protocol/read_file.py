import xlrd
import xlsxwriter
from docx import Document

from table import Table

workbook = xlsxwriter.Workbook("E:\Temp\\text.xlsx")
doc = Document('E:\Temp\\test.docx')
tableList = []

table_title_format = workbook.add_format({
    'bold': True,
    'border': 6,
    'align': 'center',  # 水平居中
    'valign': 'vcenter',  # 垂直居中
    'fg_color': '#D7E4BC',  # 颜色填充
})

# 重复表的张数
repet_table = []

for table in doc.tables:
    rows = table.rows
    length = len(rows)
    i = 0
    t = Table()
    listRow = []
    validateTable = True
    while i < length:
        cells = rows[i].cells
        if i == 0:
            tn = cells[0].text
            t.tableName = cells[0].text
            if t.tableName.find("Example:") == -1 or t.tableName.find("GT") == -1:
                validateTable = False
                break
            t.tableName = t.tableName.replace("Example:", "").lstrip(" \n")
            t.tableName = t.tableName[0:t.tableName.find('GT') + 5]
            originalName = t.tableName
            t.tableName = t.tableName.replace(":", "-")
            try:
                worksheet = workbook.add_worksheet(t.tableName)
                worksheet.merge_range('A1:D1', originalName, table_title_format)
            except Exception:
                repet_table.append(originalName)
                break
        elif i == 1:
            for colNum, c in enumerate(t.colName):
                worksheet.write(i, colNum, c)
        else:
            row = []
            for colNum, c in enumerate(cells):
                row.append(c.text)
                content = c.text if not c.text.startswith('\'') else '\''+c.text
                worksheet.write(i, colNum, c.text)
            listRow.append(row)
        i += 1
    if validateTable:
        t.body = listRow
        # t.tostring()
        tableList.append(t)

workbook.close()

print("一共处理 " + str(len(tableList)) + " 张表")
print("一共有 " + str(len(repet_table)) + " 张表产生冲突")
for name in repet_table:
    print(name)

xlsx = xlrd.open_workbook("E:\Temp\\white_list.xlsx")
sheet = xlsx.sheet_by_name("AT+GTBSI")
white_list = {}
for r in sheet._cell_values:
    white_list[r[0]] = r[1]

for r in tableList[0].body:
    if r[0] in white_list:
        if white_list[r[0]] == r[2]:
            print("匹配:" + r[0])
        else:
            print("不匹配:" + r[0])
    else:
        print("键不存在:" + r[0])
